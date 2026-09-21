from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
SCREENSHOTS = os.path.join(ROOT, "screenshots")
DIAGRAM = os.path.join(ROOT, "diagram", "usecase-diagram.png")

DARK_BG = RGBColor(0x0B, 0x0E, 0x14)
ACCENT = RGBColor(0x7D, 0xD3, 0xFC)
TEXT = RGBColor(0xE6, 0xE9, 0xEF)
SUBTEXT = RGBColor(0x8A, 0x92, 0xA6)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG


def add_title(slide, text, size=34):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.9))
    tf = box.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    return box


def add_body(slide, lines, top=1.3, size=17, left=0.8, width=11.7):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(6)
    return box


def add_image_slide(slide, title, img_path):
    add_bg(slide)
    add_title(slide, title, size=30)
    pic_w = Inches(9.5)
    left = (Inches(13.33) - pic_w) / 2
    slide.shapes.add_picture(img_path, left, Inches(1.25), width=pic_w)


# Slide 1: Title
s = prs.slides.add_slide(blank)
add_bg(s)
box = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11), Inches(1.5))
p = box.text_frame.paragraphs[0]
p.text = "Lunar Ice Watch"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = ACCENT
sub = s.shapes.add_textbox(Inches(1), Inches(3.7), Inches(11), Inches(1))
p2 = sub.text_frame.paragraphs[0]
p2.text = "South Pole Candidate Ice Site Screening Dashboard"
p2.font.size = Pt(24)
p2.font.color.rgb = TEXT
sub2 = s.shapes.add_textbox(Inches(1), Inches(4.4), Inches(11), Inches(1))
p3 = sub2.text_frame.paragraphs[0]
p3.text = "Full Stack Development Course - Mid Semester Evaluation"
p3.font.size = Pt(16)
p3.font.color.rgb = SUBTEXT

# Slide 2: Overview - problem, objective, tech stack
s = prs.slides.add_slide(blank)
add_bg(s)
add_title(s, "Overview")
add_body(s, [
    "Problem: Water ice near the Moon's south pole is hard to confirm directly -",
    "permanently shadowed regions get no sunlight, so scientists rely on indirect",
    "signals (radar, temperature, terrain) that can each be misleading alone.",
    "",
    "Objective: A dashboard that screens candidate sites by combining those",
    "signals into one clear, explainable verdict per site.",
    "",
    "Features: login, site dashboard, upload/run analysis, per-site evidence",
    "breakdown, hazard map, and a dedicated rover traverse simulation.",
    "",
    "Tech stack: HTML5 / CSS3 / vanilla JavaScript frontend, Node.js 'http'",
    "static file server, client-side JSON dataset, HTML5 Canvas for map + rover."
], size=17)

# Slide 3: Use case diagram
s = prs.slides.add_slide(blank)
add_image_slide(s, "Use Case Diagram", DIAGRAM)

# Slide 4: Screens
s = prs.slides.add_slide(blank)
add_bg(s)
add_title(s, "Screens", size=30)

screens = [
    ("Login", "01-login.jpg"),
    ("Dashboard", "02-dashboard.jpg"),
    ("Site Evidence", "03-site-detail.jpg"),
    ("Upload / Run Analysis", "04-upload.jpg"),
    ("Hazard Map", "05-map.jpg"),
    ("Rover Traverse", "06-traverse.jpg"),
]

col_w = Inches(4.04)
gap_x = Inches(0.2)
row_h = Inches(2.4)
x0 = Inches(0.4)
y0 = Inches(1.2)

for i, (title, fname) in enumerate(screens):
    row = i // 3
    col = i % 3
    x = x0 + col * (col_w + gap_x)
    y = y0 + row * row_h
    pic = s.shapes.add_picture(os.path.join(SCREENSHOTS, fname), x, y, width=col_w)
    cap = s.shapes.add_textbox(x, y + pic.height + Inches(0.03), col_w, Inches(0.3))
    cp = cap.text_frame.paragraphs[0]
    cp.text = title
    cp.font.size = Pt(12)
    cp.font.color.rgb = SUBTEXT
    cp.alignment = PP_ALIGN.CENTER

# Slide 5: Future scope + thank you
s = prs.slides.add_slide(blank)
add_bg(s)
add_title(s, "Future Scope")
add_body(s, [
    "- Replace the mock dataset with a real backend database",
    "- Add real user authentication and role-based access",
    "- Persist uploaded data files and pipeline results",
    "- Connect the evidence engine to actual scoring logic",
    "- Add 3D terrain visualization for the hazard map",
], top=1.3, size=18)

box = s.shapes.add_textbox(Inches(1), Inches(5.6), Inches(11), Inches(1.2))
p = box.text_frame.paragraphs[0]
p.text = "Thank You - Questions?"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = ACCENT
p.alignment = PP_ALIGN.CENTER

out_path = os.path.join(BASE, "Evaluation.pptx")
prs.save(out_path)
print("saved:", out_path)
